#!/usr/bin/env python3
"""Public QCEW metro-year predictive screen; not a causal policy model."""
from __future__ import annotations

import csv
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.dummy import DummyRegressor
from sklearn.ensemble import RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.linear_model import Ridge
from sklearn.metrics import mean_absolute_error, r2_score
from sklearn.model_selection import LeaveOneGroupOut
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from xgboost import XGBRegressor

ARTEM = Path(__file__).resolve().parents[1]
OUT = ARTEM / "data" / "nine_metro_models"
METRO_OUTPUT_NAME = "metro_year_qcew_2021_2025.csv"
RAW = OUT / "raw_qcew_industry_slices"
CENSUS = ARTEM / "data" / "peer_metro_regression" / "raw_inputs" / "census_cbsa_county_delineation_2023.xlsx"
METROS = {"12060": "Atlanta", "12420": "Austin", "16740": "Charlotte", "19100": "Dallas", "19740": "Denver", "26420": "Houston", "34980": "Nashville", "39580": "Raleigh", "33100": "Miami"}
INDUSTRIES = {"10": "total", "1013": "manufacturing", "51": "information", "52": "finance", "54": "professional", "62": "health", "72": "hospitality"}
YEARS = range(2021, 2026)
FEATURES = ["prior_employment_growth", "avg_weekly_wage", "establishments_per_100_jobs", "share_manufacturing", "share_information", "share_finance", "share_professional", "share_health", "share_hospitality"]
URL = "https://data.bls.gov/cew/data/api/{year}/4/industry/{industry}.csv"


def counties():
    ws = load_workbook(CENSUS, read_only=True, data_only=True).active
    data = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        cbsa = str(row[0]).zfill(5)
        if cbsa in METROS:
            data.append({"metro": METROS[cbsa], "cbsa": cbsa, "county_fips": str(row[9]).zfill(2) + str(row[10]).zfill(3), "county": row[7], "state": row[8]})
    df = pd.DataFrame(data)
    if df.duplicated("county_fips").any() or set(df.metro) != set(METROS.values()):
        raise ValueError("Census counties missing or assigned twice")
    return df


def fetch(pair):
    year, industry = pair
    path = RAW / f"qcew_{year}_q4_industry_{industry}.csv"
    if path.exists() and path.stat().st_size > 1000:
        return path
    url = URL.format(year=year, industry=industry)
    for attempt in range(3):
        try:
            data = subprocess.run(["curl", "--fail", "--location", "--silent", "--show-error", "--max-time", "45", url], check=True, capture_output=True).stdout
            if not data.startswith(b'"area_fips"'):
                raise ValueError(f"Unexpected source response: {url}")
            path.write_bytes(data)
            return path
        except Exception:
            if attempt == 2:
                raise
            time.sleep(attempt + 1)


def build_data(membership):
    county_set = set(membership.county_fips)
    items = [(y, i) for y in YEARS for i in INDUSTRIES]
    with ThreadPoolExecutor(max_workers=6) as pool:
        list(pool.map(fetch, items))
    rows = []
    for year, industry in items:
        path = RAW / f"qcew_{year}_q4_industry_{industry}.csv"
        with path.open(encoding="utf-8-sig", newline="") as f:
            for record in csv.DictReader(f):
                if record["area_fips"] not in county_set or record["own_code"] != "5" or record["industry_code"] != industry or record["size_code"] != "0":
                    continue
                rows.append({"county_fips": record["area_fips"], "year": year, "industry": INDUSTRIES[industry], "disclosure_code": record["disclosure_code"], "employment": float(record["month3_emplvl"]), "q4_employment_average": sum(float(record[f"month{x}_emplvl"]) for x in (1, 2, 3)) / 3, "quarterly_wages": float(record["total_qtrly_wages"]), "establishments": float(record["qtrly_estabs"])})
    source = pd.DataFrame(rows).merge(membership, on="county_fips", validate="many_to_one")
    if source.duplicated(["county_fips", "year", "industry"]).any():
        raise ValueError("Duplicate county-industry-year records")
    source.to_csv(OUT / "county_qcew_input_rows.csv", index=False)
    result = []
    for metro, counties_df in membership.groupby("metro"):
        for year in YEARS:
            total = source[(source.metro == metro) & (source.year == year) & (source.industry == "total")]
            if len(total) != len(counties_df) or total.disclosure_code.ne("").any():
                raise ValueError(f"Missing or suppressed total in {metro} {year}")
            emp = total.employment.sum()
            item = {"metro": metro, "year": year, "employment_q4_december": emp, "establishments_q4": total.establishments.sum(), "avg_weekly_wage": total.quarterly_wages.sum() / (13 * total.q4_employment_average.sum()), "establishments_per_100_jobs": 100 * total.establishments.sum() / emp}
            for industry in INDUSTRIES.values():
                if industry == "total":
                    continue
                subset = source[(source.metro == metro) & (source.year == year) & (source.industry == industry)]
                # BLS: an absent industry is zero; suppressed values are not zero.
                suppressed = int(subset.disclosure_code.ne("").sum())
                item[f"share_{industry}"] = np.nan if suppressed else subset.employment.sum() / emp
                item[f"suppressed_counties_{industry}"] = suppressed
            result.append(item)
    metro_df = pd.DataFrame(result).sort_values(["metro", "year"])
    metro_df["prior_employment_growth"] = metro_df.groupby("metro").employment_q4_december.pct_change(fill_method=None)
    metro_df.to_csv(OUT / METRO_OUTPUT_NAME, index=False)
    model = metro_df.copy()
    model["target_year"] = model.year + 1
    nxt = metro_df[["metro", "year", "employment_q4_december"]].rename(columns={"year": "target_year", "employment_q4_december": "next_employment"})
    model = model.merge(nxt, on=["metro", "target_year"], validate="one_to_one")
    model["target_growth"] = model.next_employment / model.employment_q4_december - 1
    model = model.dropna(subset=["prior_employment_growth", "target_growth"]).copy()
    model.to_csv(OUT / "model_panel.csv", index=False)
    return metro_df, model


def model_factory(name):
    if name == "mean_baseline":
        return DummyRegressor(strategy="mean")
    if name == "ridge":
        return make_pipeline(SimpleImputer(strategy="median", keep_empty_features=True), StandardScaler(), Ridge(alpha=10.0))
    if name == "random_forest":
        return make_pipeline(SimpleImputer(strategy="median", keep_empty_features=True), RandomForestRegressor(n_estimators=250, min_samples_leaf=3, max_features=0.8, random_state=42, n_jobs=1))
    return make_pipeline(SimpleImputer(strategy="median", keep_empty_features=True), XGBRegressor(n_estimators=100, max_depth=2, min_child_weight=5, learning_rate=0.04, subsample=1.0, colsample_bytree=0.8, reg_lambda=5.0, random_state=42, n_jobs=1))


def evaluate(panel):
    X, y = panel[FEATURES], panel.target_growth
    predictions, importance = [], []
    time_folds = [(np.flatnonzero(panel.target_year < yr), np.flatnonzero(panel.target_year == yr), str(yr)) for yr in (2024, 2025)]
    group_folds = [(tr, te, str(panel.iloc[te].metro.iloc[0])) for tr, te in LeaveOneGroupOut().split(X, y, panel.metro)]
    for split_name, folds in [("forward_year", time_folds), ("leave_one_metro_out", group_folds)]:
        for train, test, label in folds:
            for name in ("mean_baseline", "ridge", "random_forest", "xgboost"):
                estimator = model_factory(name)
                estimator.fit(X.iloc[train], y.iloc[train])
                pred = estimator.predict(X.iloc[test])
                for pos, value in zip(test, pred):
                    row = panel.iloc[pos]
                    predictions.append({"evaluation": split_name, "held_out": label, "model": name, "metro": row.metro, "target_year": row.target_year, "observed_growth": row.target_growth, "predicted_growth": float(value)})
                if split_name == "forward_year" and name in ("random_forest", "xgboost"):
                    baseline = mean_absolute_error(y.iloc[test], pred)
                    rng = np.random.default_rng(101)
                    for feature in FEATURES:
                        changes = []
                        for _ in range(100):
                            shuffled = X.iloc[test].copy()
                            shuffled[feature] = rng.permutation(shuffled[feature].to_numpy())
                            changes.append(100 * (mean_absolute_error(y.iloc[test], estimator.predict(shuffled)) - baseline))
                        importance.append({"model": name, "held_out_year": label, "feature": feature, "permutation_delta_mae_pp": float(np.mean(changes)), "permutation_std_pp": float(np.std(changes))})
    pred_df = pd.DataFrame(predictions)
    pred_df.to_csv(OUT / "held_out_predictions.csv", index=False)
    metrics = []
    for (evaluation, name), group in pred_df.groupby(["evaluation", "model"]):
        metrics.append({"evaluation": evaluation, "model": name, "n": len(group), "mae_percentage_points": 100 * mean_absolute_error(group.observed_growth, group.predicted_growth), "pooled_r2": r2_score(group.observed_growth, group.predicted_growth)})
    pd.DataFrame(metrics).to_csv(OUT / "model_comparison.csv", index=False)
    pd.DataFrame(importance).to_csv(OUT / "heldout_permutation_importance.csv", index=False)
    return pd.DataFrame(metrics), pd.DataFrame(importance)


def main():
    RAW.mkdir(parents=True, exist_ok=True)
    membership = counties()
    membership.to_csv(OUT / "metro_counties_2023_census.csv", index=False)
    metro, panel = build_data(membership)
    metrics, importance = evaluate(panel)
    print("Counties:", membership.groupby("metro").size().to_dict())
    print("Model rows:", len(panel), "outcome years:", sorted(panel.target_year.unique()))
    print("Missing predictors:", panel[FEATURES].isna().sum().to_dict())
    print(metrics.to_string(index=False))
    print("Mean held-out importance (percentage-point MAE):")
    print(importance.groupby(["model", "feature"]).permutation_delta_mae_pp.mean().sort_values(ascending=False).to_string())


if __name__ == "__main__":
    main()
